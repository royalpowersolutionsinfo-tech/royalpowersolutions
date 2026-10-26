import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_quotation_docx():
    doc = docx.Document()
    
    # Page setup - Margins
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
    PRIMARY = RGBColor(30, 58, 138)     # Navy #1E3A8A
    SECONDARY = RGBColor(14, 116, 144)  # Cyan #0E7490
    DARK = RGBColor(31, 41, 55)         # Charcoal #1F2937
    GRAY = RGBColor(71, 85, 105)        # Muted #475569
    
    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("COMMERCIAL QUOTATION & SCOPE OF WORK")
    run_title.font.name = "Segoe UI"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = PRIMARY
    
    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run("Complete Digital Presence, Local SEO Engine, Branding & 1-Year Maintenance Package")
    run_sub.font.name = "Segoe UI"
    run_sub.font.size = Pt(10.5)
    run_sub.font.italic = True
    run_sub.font.color.rgb = SECONDARY
    
    # Info Table
    info_table = doc.add_table(rows=2, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.autofit = False
    
    col_widths = [Inches(3.5), Inches(3.5)]
    for row in info_table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width
            
    # Headers
    h0 = info_table.cell(0, 0).paragraphs[0]
    r = h0.add_run("QUOTATION DETAILS")
    r.font.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)
    r.font.size = Pt(9.5)
    set_cell_background(info_table.cell(0, 0), "1E3A8A")
    set_cell_margins(info_table.cell(0, 0), top=120, bottom=120, left=150, right=150)
    
    h1 = info_table.cell(0, 1).paragraphs[0]
    r = h1.add_run("CLIENT INFORMATION")
    r.font.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)
    r.font.size = Pt(9.5)
    set_cell_background(info_table.cell(0, 1), "1E3A8A")
    set_cell_margins(info_table.cell(0, 1), top=120, bottom=120, left=150, right=150)
    
    # Content Left
    c_left = info_table.cell(1, 0)
    set_cell_background(c_left, "EFF6FF")
    set_cell_margins(c_left, top=140, bottom=140, left=150, right=150)
    p = c_left.paragraphs[0]
    p.paragraph_format.line_spacing = 1.3
    
    r = p.add_run("Quote Ref: ")
    r.font.bold = True
    p.add_run("RPS-QTN-2026-001-REV1\n")
    r = p.add_run("Date: ")
    r.font.bold = True
    p.add_run("September 29, 2026\n")
    r = p.add_run("Validity: ")
    r.font.bold = True
    p.add_run("30 Days\n")
    r = p.add_run("Currency: ")
    r.font.bold = True
    p.add_run("INR (₹)")
    for run in p.runs:
        run.font.name = "Segoe UI"
        run.font.size = Pt(9)
        run.font.color.rgb = DARK
        
    # Content Right
    c_right = info_table.cell(1, 1)
    set_cell_background(c_right, "EFF6FF")
    set_cell_margins(c_right, top=140, bottom=140, left=150, right=150)
    p2 = c_right.paragraphs[0]
    p2.paragraph_format.line_spacing = 1.3
    
    r = p2.add_run("Client: ")
    r.font.bold = True
    p2.add_run("Royal Power Solutions\n")
    r = p2.add_run("Location: ")
    r.font.bold = True
    p2.add_run("Chennai, Tamil Nadu\n")
    r = p2.add_run("Scope: ")
    r.font.bold = True
    p2.add_run("Website, 120 Local Pages, Logo, Google Profile, 1-Yr Maint.\n")
    r = p2.add_run("Target Market: ")
    r.font.bold = True
    p2.add_run("Chennai & Surrounding Districts")
    for run in p2.runs:
        run.font.name = "Segoe UI"
        run.font.size = Pt(9)
        run.font.color.rgb = DARK

    # Spacer
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(8)
    p_sp.paragraph_format.space_after = Pt(2)
    
    # 4 Badges Overview Grid
    badge_table = doc.add_table(rows=1, cols=4)
    badge_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    badge_widths = [Inches(1.75), Inches(1.75), Inches(1.75), Inches(1.75)]
    
    badges_data = [
        ("158", "TOTAL LIVE PAGES"),
        ("66", "TARGET KEYWORDS"),
        ("1 YEAR", "HOSTING & MAINT."),
        ("₹ 50,000", "TOTAL PACKAGE")
    ]
    
    for idx, (num_text, lbl_text) in enumerate(badges_data):
        cell = badge_table.cell(0, idx)
        cell.width = badge_widths[idx]
        set_cell_background(cell, "F8FAFC")
        set_cell_margins(cell, top=120, bottom=120, left=100, right=100)
        p_b = cell.paragraphs[0]
        p_b.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        r1 = p_b.add_run(num_text + "\n")
        r1.font.name = "Segoe UI"
        r1.font.size = Pt(15)
        r1.font.bold = True
        r1.font.color.rgb = PRIMARY
        
        r2 = p_b.add_run(lbl_text)
        r2.font.name = "Segoe UI"
        r2.font.size = Pt(7.5)
        r2.font.bold = True
        r2.font.color.rgb = GRAY
        
    # Section 1 Heading
    p_s1 = doc.add_paragraph()
    p_s1.paragraph_format.space_before = Pt(16)
    p_s1.paragraph_format.space_after = Pt(6)
    r = p_s1.add_run("1. Executive Summary & Deliverables Breakdown (Simple & Detailed)")
    r.font.name = "Segoe UI"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = PRIMARY

    bullets = [
        ("Total Pages Built & Indexed (158 Pages Total): ", [
            "• 8 Core Business Pages: Home, About Us, Services, Projects Showcase, Gallery, Blog, Contact, and Custom 404.",
            "• 120 Programmatic Local SEO Landing Pages: Covering 6 major Chennai commercial hubs (Guindy, Saidapet, West Mambalam, Ashok Pillar, Jafferkhanpet, Chennai City) across 20 high-demand power services (6 hubs × 20 services = 120 geo-targeted pages).",
            "• 20 Dedicated Service & Product Pages: Specifications for 1KVA to 1000KVA Online UPS, Batteries, Stabilizers, Transformers, AMC, and Solar.",
            "• 10 Commercial Project Case Studies & Technical Blog Guides driving qualified organic inquiries."
        ]),
        ("Target Search Keywords (66 Strategically Researched Keywords): ", [
            "• 5 Brand Identity Keywords (Royal Power Solutions, etc.)",
            "• 18 High-Priority Commercial Money Keywords (Online UPS sales Chennai, UPS repair service near me, etc.)",
            "• 21 Local Area Geo-Targeted Keywords (Guindy, Saidapet, Ashok Pillar, West Mambalam, Jafferkhanpet)",
            "• 14 Industrial & B2B Keywords (Three-phase UPS AMC, isolation transformer installation, industrial stabilizer rentals)",
            "• 8 Informational & Blog Intent Keywords (Troubleshooting, maintenance guides, error codes)"
        ]),
        ("Brand Identity & Logo Design (Included): ", [
            "• Custom modern vector logo designed specifically for Royal Power Solutions.",
            "• High-resolution vector master files (SVG, PDF) and web formats (transparent PNG, JPEG).",
            "• Website Favicon, brand typography, and official primary color palette setup."
        ]),
        ("Google Business Profile (GBP / GMB) Creation & Local Optimization (Included): ", [
            "• Official Google Profile creation, NAP verification, and Google Maps pin calibration.",
            "• Categories selection (UPS Supplier, Electrical Contractor, Battery Store), business hours & Chennai service areas.",
            "• Geo-tagged photos upload, service catalog sync, and 1-click Google Review direct link setup."
        ]),
        ("Cloud Hosting & 1-Year Technical Maintenance AMC (Included): ", [
            "• High-speed cloud server deployment with domain DNS configuration and lifetime HTTPS / SSL certificate.",
            "• 1 Full Year Technical Maintenance: Automated monthly cloud backups, 99.9% uptime monitoring, security patching, bug fixes, and minor content updates for 12 continuous months."
        ]),
        ("Lead Generation & Direct Customer Contact (Included): ", [
            "• 1-click floating WhatsApp chat trigger active across all 158 pages.",
            "• Emergency click-to-call direct phone dialers for immediate customer phone inquiries.",
            "• Interactive quotation request forms with instant notification backend."
        ])
    ]

    for title, items in bullets:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(2)
        r_bold = p.add_run(title)
        r_bold.font.name = "Segoe UI"
        r_bold.font.size = Pt(9.5)
        r_bold.font.bold = True
        r_bold.font.color.rgb = DARK
        
        for item in items:
            p_sub = doc.add_paragraph()
            p_sub.paragraph_format.left_indent = Inches(0.25)
            p_sub.paragraph_format.space_before = Pt(1)
            p_sub.paragraph_format.space_after = Pt(1)
            r_item = p_sub.add_run(item)
            r_item.font.name = "Segoe UI"
            r_item.font.size = Pt(8.5)
            r_item.font.color.rgb = DARK

    # Section 2 Heading
    p_s2 = doc.add_paragraph()
    p_s2.paragraph_format.space_before = Pt(16)
    p_s2.paragraph_format.space_after = Pt(6)
    r = p_s2.add_run("2. Detailed Itemized Quotation (Editable)")
    r.font.name = "Segoe UI"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = PRIMARY
    
    # Quotation Table (Sum = ₹ 50,000)
    q_table = doc.add_table(rows=10, cols=4)
    q_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    q_widths = [Inches(0.4), Inches(2.1), Inches(3.3), Inches(1.2)]
    
    headers = ["#", "Scope / Deliverable", "Detailed Specifications & Inclusions", "Amount (₹)"]
    for idx, name in enumerate(headers):
        cell = q_table.cell(0, idx)
        cell.width = q_widths[idx]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        if idx == 3:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run(name)
        r.font.name = "Segoe UI"
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    items = [
        ("1", "Professional Logo Design & Brand Identity", 
         "• Custom modern logo creation representing power, reliability & expertise\n• Vector files (SVG, PDF) + transparent PNG/JPG for web & print\n• Web browser favicon icon and primary brand color styling guide", "₹ 4,000"),
         
        ("2", "Google Business Profile (GBP / GMB) Creation & Setup", 
         "• Google Business Profile setup, category mapping & verification guidance\n• Exact Google Maps pin localization across Chennai target service areas\n• Photos upload, service catalog listing & 1-click Google Review direct link", "₹ 3,000"),
         
        ("3", "Core Website UI/UX Design & Frontend Development (8 Pages)", 
         "• Responsive, mobile-first design across smartphones, tablets & desktops\n• Sub-second speed architecture (HTML5, CSS3, Vanilla JS)\n• 8 Core Pages: Home, About Us, Services, Projects, Gallery, Blog, Contact, 404\n• High-conversion layout with sticky header, trust badges & hero callouts", "₹ 12,000"),
         
        ("4", "Service & Product Catalog Showcase Module", 
         "• 20 dedicated service & product pages detailing technical specifications\n• Full coverage: 1KVA–1000KVA Online UPS, Batteries, Stabilizers, Transformers\n• Interactive completed projects portfolio showcase and photo gallery\n• Partner brand trust badges (Numeric, APC, Emerson/Vertiv, Exide, Amaron)", "₹ 6,000"),
         
        ("5", "Programmatic Local SEO Engine (120 Chennai Area Pages)", 
         "• Automated local search landing page engine targeting 6 prime Chennai hubs:\n  (Guindy, Saidapet, West Mambalam, Ashok Pillar, Jafferkhanpet, Chennai City)\n• 20 dedicated service landing pages per locality (Total 120 Geo-targeted pages)\n• Area-specific H1 titles, customized local intros & neighborhood service FAQs", "₹ 10,000"),
         
        ("6", "SEO Strategy, Keyword Research (66 Keywords) & Technical SEO", 
         "• 66 high-intent keywords across Brand, Money, Geo Local, B2B & Blog categories\n• XML Sitemap with 158 live indexable URLs submitted to Google Search Console\n• Optimized robots.txt crawler directives and canonical tags\n• Schema.org structured data (LocalBusiness, Service, Breadcrumbs, Star Ratings)", "₹ 5,000"),
         
        ("7", "Lead Generation & Instant Conversion Tools", 
         "• Floating WhatsApp 1-click chat trigger button active across all 158 pages\n• Click-to-call emergency hotline buttons for direct phone call generation\n• Interactive quotation request forms and inquiry capture with notification API\n• Embedded interactive Google Maps for local shop trust and direct navigation", "₹ 3,000"),
         
        ("8", "1-Year Cloud Hosting & Complete Maintenance (12 Months Included)", 
         "• 1 Full Year high-speed cloud hosting server deployment & domain DNS setup\n• Lifetime HTTPS / SSL security certificate integration\n• 1 Full Year Technical AMC & Maintenance: Regular monthly backups, uptime monitoring, security patching, bug fixes & minor content updates for 12 months", "₹ 7,000"),
    ]
    
    for row_idx, (num, scope, desc, price) in enumerate(items, start=1):
        row = q_table.rows[row_idx]
        for col_idx, width in enumerate(q_widths):
            row.cells[col_idx].width = width
            set_cell_margins(row.cells[col_idx], top=80, bottom=80, left=90, right=90)
            if row_idx % 2 == 0:
                set_cell_background(row.cells[col_idx], "F8FAFC")
                
        # Cell 0: Num
        p = row.cells[0].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(num)
        r.font.name = "Segoe UI"
        r.font.size = Pt(8.5)
        r.font.bold = True
        
        # Cell 1: Scope
        p = row.cells[1].paragraphs[0]
        r = p.add_run(scope)
        r.font.name = "Segoe UI"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = PRIMARY
        
        # Cell 2: Desc
        p = row.cells[2].paragraphs[0]
        r = p.add_run(desc)
        r.font.name = "Segoe UI"
        r.font.size = Pt(8)
        r.font.color.rgb = DARK
        
        # Cell 3: Price
        p = row.cells[3].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run(price)
        r.font.name = "Segoe UI"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = DARK

    # Total Row (Row 9)
    total_row = q_table.rows[9]
    for c_idx in range(4):
        total_row.cells[c_idx].width = q_widths[c_idx]
        set_cell_background(total_row.cells[c_idx], "0E7490")
        set_cell_margins(total_row.cells[c_idx], top=120, bottom=120, left=90, right=90)
        
    # Merge cells 0, 1, 2 for total label
    total_row.cells[0].merge(total_row.cells[1]).merge(total_row.cells[2])
    p_tot_lbl = total_row.cells[0].paragraphs[0]
    p_tot_lbl.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p_tot_lbl.add_run("TOTAL ALL-INCLUSIVE PACKAGE AMOUNT:  ")
    r.font.name = "Segoe UI"
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)
    
    p_tot_val = total_row.cells[3].paragraphs[0]
    p_tot_val.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p_tot_val.add_run("₹ 50,000")
    r.font.name = "Segoe UI"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)
    
    # Notice Box
    p_box = doc.add_paragraph()
    p_box.paragraph_format.space_before = Pt(8)
    p_box.paragraph_format.space_after = Pt(10)
    r_n = p_box.add_run("✓ ALL-INCLUSIVE PACKAGE GUARANTEE: ")
    r_n.font.name = "Segoe UI"
    r_n.font.bold = True
    r_n.font.size = Pt(8.5)
    r_n.font.color.rgb = RGBColor(5, 150, 105)
    r_n2 = p_box.add_run("This quotation includes complete web development, 158 pages, 66 keywords, logo design, Google Business Profile creation, and 1 full year of cloud hosting & technical maintenance for a fixed total of ")
    r_n2.font.name = "Segoe UI"
    r_n2.font.size = Pt(8.5)
    r_n3 = p_box.add_run("₹ 50,000 only")
    r_n3.font.name = "Segoe UI"
    r_n3.font.bold = True
    r_n3.font.size = Pt(8.5)
    r_n4 = p_box.add_run(". There are no hidden or surprise charges.")
    r_n4.font.name = "Segoe UI"
    r_n4.font.size = Pt(8.5)

    # Section 3: Ongoing Renewal Terms
    p_s3 = doc.add_paragraph()
    p_s3.paragraph_format.space_before = Pt(12)
    p_s3.paragraph_format.space_after = Pt(4)
    r = p_s3.add_run("3. Ongoing Renewal Terms (Applicable After 1st Year)")
    r.font.name = "Segoe UI"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = PRIMARY
    
    rn_table = doc.add_table(rows=3, cols=5)
    rn_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    rn_widths = [Inches(0.4), Inches(2.3), Inches(2.8), Inches(0.8), Inches(1.1)]
    
    rn_headers = ["#", "Service Name", "Scope of Support", "Period", "Amount (₹)"]
    for idx, name in enumerate(rn_headers):
        cell = rn_table.cell(0, idx)
        cell.width = rn_widths[idx]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        p = cell.paragraphs[0]
        if idx == 4:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run(name)
        r.font.name = "Segoe UI"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    rn_items = [
        ("A", "Annual Hosting, SSL & Technical Maintenance (Year 2 Onwards)", 
         "Cloud server renewal, domain renewal support, SSL maintenance, monthly backups, 99.9% uptime monitoring, and security patching. (Note: Year 1 is 100% FREE & INCLUDED in the ₹ 50,000 package).", "Annual", "₹ 8,000 / yr"),
        ("B", "Monthly Local SEO & Keyword Rank Tracking (Optional)", 
         "Monthly keyword ranking audit for all 66 keywords, Google Search Console monitoring, Google Business Profile monthly updates, and 2 fresh SEO blog articles monthly.", "Monthly", "₹ 7,500 / mo")
    ]
    
    for row_idx, (num, sname, sdesc, period, pr) in enumerate(rn_items, start=1):
        row = rn_table.rows[row_idx]
        for col_idx, width in enumerate(rn_widths):
            row.cells[col_idx].width = width
            set_cell_margins(row.cells[col_idx], top=70, bottom=70, left=80, right=80)
            
        p = row.cells[0].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(num)
        r.font.name = "Segoe UI"
        r.font.size = Pt(8.5)
        r.font.bold = True
        
        p = row.cells[1].paragraphs[0]
        r = p.add_run(sname)
        r.font.name = "Segoe UI"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = PRIMARY
        
        p = row.cells[2].paragraphs[0]
        r = p.add_run(sdesc)
        r.font.name = "Segoe UI"
        r.font.size = Pt(8)
        
        p = row.cells[3].paragraphs[0]
        r = p.add_run(period)
        r.font.name = "Segoe UI"
        r.font.size = Pt(8)
        
        p = row.cells[4].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run(pr)
        r.font.name = "Segoe UI"
        r.font.size = Pt(8.5)
        r.font.bold = True

    # Section 4: Payment Terms & Milestones
    p_s4 = doc.add_paragraph()
    p_s4.paragraph_format.space_before = Pt(14)
    p_s4.paragraph_format.space_after = Pt(4)
    r = p_s4.add_run("4. Payment Terms & Milestones")
    r.font.name = "Segoe UI"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = PRIMARY
    
    m_list = [
        ("Milestone 1 — Advance Booking / Project Initiation (40%): ", "₹ 20,000 — Payable upon contract sign-off to initiate logo design, domain configuration, and website structure."),
        ("Milestone 2 — Design, Core Pages & Google Profile Approval (30%): ", "₹ 15,000 — Payable upon review and approval of custom logo, core 8 pages, and Google Business Profile setup."),
        ("Milestone 3 — Final Delivery, 120 Local SEO Pages & 1-Year Live Handover (30%): ", "₹ 15,000 — Payable upon completion of all 120 programmatic pages, technical SEO audit, and 1-year live hosting handover.")
    ]
    
    for m_title, m_desc in m_list:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run("• " + m_title)
        r1.font.name = "Segoe UI"
        r1.font.size = Pt(8.5)
        r1.font.bold = True
        r1.font.color.rgb = DARK
        r2 = p.add_run(m_desc)
        r2.font.name = "Segoe UI"
        r2.font.size = Pt(8.5)
        r2.font.color.rgb = DARK

    # Section 5: Payment Details & Authorization
    p_s5 = doc.add_paragraph()
    p_s5.paragraph_format.space_before = Pt(14)
    p_s5.paragraph_format.space_after = Pt(4)
    r = p_s5.add_run("5. Payment Details & Client Acceptance")
    r.font.name = "Segoe UI"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = PRIMARY
    
    sign_table = doc.add_table(rows=1, cols=2)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sign_widths = [Inches(3.5), Inches(3.5)]
    
    c_pay = sign_table.cell(0, 0)
    c_pay.width = sign_widths[0]
    set_cell_background(c_pay, "EFF6FF")
    set_cell_margins(c_pay, top=120, bottom=120, left=140, right=140)
    p = c_pay.paragraphs[0]
    p.paragraph_format.line_spacing = 1.3
    r = p.add_run("PAYMENT DETAILS (NEFT / UPI / IMPS)\n\n")
    r.font.bold = True
    r.font.color.rgb = PRIMARY
    p.add_run("Account Name: [Your Name / Company Name]\n")
    p.add_run("Bank Name: [Your Bank Name]\n")
    p.add_run("Account Number: [Enter Account Number]\n")
    p.add_run("IFSC Code: [Enter IFSC Code]\n")
    p.add_run("UPI ID / GPay / PhonePe: [Enter Phone/UPI]")
    for run in p.runs:
        run.font.name = "Segoe UI"
        run.font.size = Pt(8.5)
        if not run.font.color.rgb:
            run.font.color.rgb = DARK
            
    c_sign = sign_table.cell(0, 1)
    c_sign.width = sign_widths[1]
    set_cell_background(c_sign, "F8FAFC")
    set_cell_margins(c_sign, top=120, bottom=120, left=140, right=140)
    p = c_sign.paragraphs[0]
    p.paragraph_format.line_spacing = 1.3
    r = p.add_run("CLIENT ACCEPTANCE & SIGNATURE\n\n")
    r.font.bold = True
    r.font.color.rgb = PRIMARY
    p.add_run("Accepted By: __________________________\n\n")
    p.add_run("Designation: __________________________\n\n")
    p.add_run("Signature: ____________________________\n\n")
    p.add_run("Date: _________________________________")
    for run in p.runs:
        run.font.name = "Segoe UI"
        run.font.size = Pt(8.5)
        if not run.font.color.rgb:
            run.font.color.rgb = DARK

    doc.save("Royal_Power_Solutions_Quotation.docx")
    print("Royal_Power_Solutions_Quotation.docx created successfully!")

if __name__ == "__main__":
    create_quotation_docx()
